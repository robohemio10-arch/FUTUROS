"""Fail-closed read-only loader for the official post-OCR trades master."""

from __future__ import annotations

import hashlib
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .contracts import (
    CANONICAL_SOURCE_CONTRACT,
    MasterSourceContract,
    contract_sha256,
)


class OfficialMasterValidationError(ValueError):
    """Raised when the official master violates its frozen source contract."""

    def __init__(
        self,
        code: str,
        *,
        details: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(code)
        self.code = code
        self.details = details or {}


@dataclass(frozen=True)
class SourceAudit:
    status: str
    master_path: str
    master_sha256: str
    source_contract_sha256: str
    rows: int
    columns: int
    sheets: tuple[str, ...]
    trades_dimension: str
    formula_count: int
    column_names: tuple[str, ...]
    economic_identity_max_abs_error: float
    fee_regime_counts: dict[str, int]
    invalid_order_id_count: int
    trade_sequence_contiguous: bool
    hash_unchanged_after_read: bool
    read_only: bool
    writes_master: bool
    sends_orders: bool
    changes_risk: bool
    operational_authority: bool

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class OfficialMasterData:
    path: Path
    frame: pd.DataFrame
    audit: SourceAudit


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def _fail(
    code: str,
    **details: Any,
) -> None:
    raise OfficialMasterValidationError(code, details=details)


def _numeric_series(
    frame: pd.DataFrame,
    column: str,
) -> pd.Series:
    if column not in frame.columns:
        _fail("missing_required_column", column=column)

    series = pd.to_numeric(frame[column], errors="coerce")

    if series.isna().any():
        _fail(
            "non_numeric_or_missing_values",
            column=column,
            invalid_count=int(series.isna().sum()),
        )

    values = series.to_numpy(dtype=float)
    invalid_finite = int((~np.isfinite(values)).sum())

    if invalid_finite:
        _fail(
            "non_finite_values",
            column=column,
            invalid_count=invalid_finite,
        )

    return series.astype(float)


def _read_trades_sheet(
    worksheet: Any,
    expected_columns: int,
) -> tuple[list[str], list[tuple[Any, ...]]]:
    rows = list(
        worksheet.iter_rows(
            min_row=1,
            max_row=worksheet.max_row,
            max_col=expected_columns,
            values_only=True,
        )
    )

    if not rows:
        _fail("trades_sheet_empty")

    headers = [
        "" if value is None else str(value).strip()
        for value in rows[0]
    ]
    data = [
        tuple(row[:expected_columns])
        for row in rows[1:]
    ]
    return headers, data


def _count_formulas(workbook: Any) -> int:
    count = 0
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows():
            for cell in row:
                if getattr(cell, "data_type", None) == "f":
                    count += 1
    return count


def _validate_trade_sequence(frame: pd.DataFrame) -> None:
    sequence = _numeric_series(frame, "trade_sequence")
    values = sequence.to_numpy(dtype=float)

    if not np.allclose(values, np.round(values), atol=0.0, rtol=0.0):
        _fail("trade_sequence_not_integer")

    expected: np.ndarray = np.arange(1, len(frame) + 1, dtype=np.int64)
    actual = np.round(values).astype(np.int64)

    if not np.array_equal(actual, expected):
        _fail(
            "trade_sequence_not_contiguous",
            first_actual=int(actual[0]) if len(actual) else None,
            last_actual=int(actual[-1]) if len(actual) else None,
        )


def _validate_order_ids(
    frame: pd.DataFrame,
    pattern: str,
) -> int:
    if "order_id_raw" not in frame.columns:
        _fail("missing_required_column", column="order_id_raw")

    order_ids = frame["order_id_raw"].astype("string").str.strip()
    matcher = re.compile(pattern)
    invalid_mask = ~order_ids.map(
        lambda value: bool(matcher.fullmatch(str(value)))
        if not pd.isna(value)
        else False
    )
    invalid_count = int(invalid_mask.sum())

    if invalid_count:
        _fail(
            "invalid_or_missing_order_id",
            invalid_count=invalid_count,
        )

    return invalid_count


def _validate_economic_identity(
    frame: pd.DataFrame,
    tolerance: float,
) -> float:
    reported = _numeric_series(frame, "reported_pnl")
    adjustment = _numeric_series(frame, "economic_fee_adjustment")
    economic = _numeric_series(frame, "economic_net_pnl")

    errors = (economic - (reported + adjustment)).abs()
    max_error = float(errors.max()) if len(errors) else 0.0

    if max_error > tolerance:
        _fail(
            "economic_identity_drift",
            max_abs_error=max_error,
            tolerance=tolerance,
        )

    return max_error


def load_official_trades_master(
    path: str | Path,
    *,
    contract: MasterSourceContract = CANONICAL_SOURCE_CONTRACT,
) -> OfficialMasterData:
    target = Path(path)

    if not target.exists() or not target.is_file():
        _fail("master_not_found", path=str(target))

    if target.is_symlink():
        _fail("master_symlink_rejected", path=str(target))

    if target.suffix.lower() != ".xlsx":
        _fail(
            "master_extension_invalid",
            suffix=target.suffix.lower(),
        )

    sha_before = sha256_file(target)

    if sha_before != contract.expected_sha256:
        _fail(
            "master_sha256_mismatch",
            expected=contract.expected_sha256,
            actual=sha_before,
        )

    workbook = load_workbook(
        target,
        read_only=True,
        data_only=False,
    )

    try:
        sheets = tuple(workbook.sheetnames)

        if sheets != contract.expected_sheets:
            _fail(
                "sheet_contract_mismatch",
                expected=list(contract.expected_sheets),
                actual=list(sheets),
            )

        trades = workbook["TRADES"]
        dimension = trades.calculate_dimension()

        if dimension != contract.expected_trades_dimension:
            _fail(
                "trades_dimension_mismatch",
                expected=contract.expected_trades_dimension,
                actual=dimension,
            )

        if trades.max_column != contract.expected_columns:
            _fail(
                "column_count_mismatch",
                expected=contract.expected_columns,
                actual=trades.max_column,
            )

        headers, data = _read_trades_sheet(
            trades,
            contract.expected_columns,
        )
        formula_count = _count_formulas(workbook)
    finally:
        workbook.close()

    sha_after = sha256_file(target)

    if sha_after != sha_before:
        _fail(
            "master_changed_during_read",
            before=sha_before,
            after=sha_after,
        )

    if formula_count != contract.expected_formula_count:
        _fail(
            "formula_count_mismatch",
            expected=contract.expected_formula_count,
            actual=formula_count,
        )

    actual_columns = tuple(headers)

    if actual_columns != contract.expected_column_order:
        _fail(
            "schema_column_order_mismatch",
            expected=list(contract.expected_column_order),
            actual=list(actual_columns),
        )

    frame = pd.DataFrame(
        data,
        columns=actual_columns,
    )

    if len(frame) != contract.expected_rows:
        _fail(
            "row_count_mismatch",
            expected=contract.expected_rows,
            actual=len(frame),
        )

    for column in contract.required_traceability_columns:
        if column not in frame.columns:
            _fail(
                "missing_traceability_column",
                column=column,
            )

        if frame[column].isna().any():
            _fail(
                "traceability_column_contains_missing",
                column=column,
                missing_count=int(frame[column].isna().sum()),
            )

    _validate_trade_sequence(frame)

    invalid_order_id_count = _validate_order_ids(
        frame,
        contract.order_id_pattern,
    )

    economic_error = _validate_economic_identity(
        frame,
        contract.economic_identity_abs_tolerance,
    )

    if "fee_semantics_regime" not in frame.columns:
        _fail(
            "missing_required_column",
            column="fee_semantics_regime",
        )

    fee_regimes = (
        frame["fee_semantics_regime"]
        .astype("string")
        .str.strip()
    )

    if fee_regimes.isna().any() or fee_regimes.eq("").any():
        _fail("fee_regime_missing_values")

    actual_regimes = set(fee_regimes.unique().tolist())
    allowed_regimes = set(contract.allowed_fee_regimes)

    if actual_regimes != allowed_regimes:
        _fail(
            "fee_regime_contract_mismatch",
            expected=sorted(allowed_regimes),
            actual=sorted(actual_regimes),
        )

    fee_counts = {
        str(key): int(value)
        for key, value in fee_regimes.value_counts().sort_index().items()
    }

    audit = SourceAudit(
        status="ok",
        master_path=str(target),
        master_sha256=sha_before,
        source_contract_sha256=contract_sha256(contract),
        rows=int(len(frame)),
        columns=int(len(frame.columns)),
        sheets=sheets,
        trades_dimension=contract.expected_trades_dimension,
        formula_count=int(formula_count),
        column_names=actual_columns,
        economic_identity_max_abs_error=economic_error,
        fee_regime_counts=fee_counts,
        invalid_order_id_count=invalid_order_id_count,
        trade_sequence_contiguous=True,
        hash_unchanged_after_read=True,
        read_only=True,
        writes_master=False,
        sends_orders=False,
        changes_risk=False,
        operational_authority=False,
    )

    return OfficialMasterData(
        path=target,
        frame=frame,
        audit=audit,
    )


__all__ = [
    "OfficialMasterData",
    "OfficialMasterValidationError",
    "SourceAudit",
    "load_official_trades_master",
    "sha256_file",
]
