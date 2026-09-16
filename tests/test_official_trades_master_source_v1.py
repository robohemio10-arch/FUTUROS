from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from smartcrypto.research.trades_master_official.contracts import (
    CANONICAL_SOURCE_CONTRACT,
    OFFICIAL_MASTER_COLUMNS,
)
from smartcrypto.research.trades_master_official.loader import (
    OfficialMasterValidationError,
    load_official_trades_master,
    sha256_file,
)


def _row(
    sequence: int,
    *,
    regime: str = "STANDARD",
    order_id: str | None = None,
    reported_pnl: float = 10.0,
    fee_adjustment: float = -1.0,
) -> dict[str, object]:
    economic = reported_pnl + fee_adjustment

    row: dict[str, object] = {
        column: ""
        for column in OFFICIAL_MASTER_COLUMNS
    }

    row.update(
        {
            "master_schema_version": "TEST_SCHEMA",
            "master_candidate_status": "PREWRITE_VALIDATED",
            "trade_sequence": sequence,
            "image_number": sequence,
            "source_file": f"source-{sequence}.jpg",
            "pnl_fechado": reported_pnl,
            "taxa_lucros_perdas_fechados_pct": 1.0,
            "preco_abertura": 100.0,
            "preco_fechamento": 101.0,
            "volume_posicao": 1.0,
            "volume_fechado": 1.0,
            "horario_abertura": (
                "SOURCE_NOT_AVAILABLE_FROM_PRINT"
                if regime == "LEGACY"
                else "2026-01-01 00:00:00"
            ),
            "horario_fechamento": "2026-01-01 00:10:00",
            "taxa_1": fee_adjustment,
            "order_id_raw": (
                order_id
                or f"{sequence:024x}"
            ),
            "moeda": "BTC/USDT",
            "fechar_side": "LONG",
            "symbol_raw": "BTC/USDT",
            "symbol": "BTC_USDT",
            "side_raw": "LONG",
            "side": "LONG",
            "classification": "MATCH_FEE_EXCLUDED",
            "fee_semantics_regime": regime,
            "reported_pnl": reported_pnl,
            "economic_fee_adjustment": fee_adjustment,
            "economic_net_pnl": economic,
            "fee_policy_version": "TEST_FEE_V1",
            "fee_policy_formula": "reported_pnl + taxa_1",
            "fee_semantics_provenance": "TEST",
            "fee_semantics_gate": "PASS",
            "economic_pnl_canonicalized": True,
            "cumulative_economic_net_pnl": economic,
            "running_peak_economic_net_pnl": economic,
            "economic_drawdown_pnl": 0.0,
            "final_dedup_group_id": "",
            "final_dedup_classification": "UNIQUE_NO_COLLISION",
            "final_dedup_canonical_image": "",
            "final_dedup_keep": True,
            "final_dedup_excluded": False,
            "final_dedup_gate": "PASS",
            "trades_master_write_allowed": False,
        }
    )

    return row


def _write_workbook(
    path: Path,
    rows: list[dict[str, object]],
) -> None:
    workbook = Workbook()
    trades = workbook.active
    trades.title = "TRADES"
    trades.append(list(OFFICIAL_MASTER_COLUMNS))

    for row in rows:
        trades.append(
            [
                row[column]
                for column in OFFICIAL_MASTER_COLUMNS
            ]
        )

    metadata = workbook.create_sheet("METADATA")
    metadata.append(["key", "value"])
    metadata.append(["test", "true"])

    audit = workbook.create_sheet("AUDIT")
    audit.append(["gate", "status"])
    audit.append(["test", "PASS"])

    workbook.save(path)
    workbook.close()


def _contract_for(
    path: Path,
    row_count: int,
):
    return replace(
        CANONICAL_SOURCE_CONTRACT,
        expected_sha256=sha256_file(path),
        expected_rows=row_count,
        expected_trades_dimension=(
            f"A1:AO{row_count + 1}"
        ),
    )


def test_loads_master_readonly_and_preserves_sha(
    tmp_path: Path,
) -> None:
    path = tmp_path / "trades_master.xlsx"
    _write_workbook(
        path,
        [
            _row(1),
            _row(2, regime="LEGACY"),
        ],
    )

    before = sha256_file(path)
    master = load_official_trades_master(
        path,
        contract=_contract_for(path, 2),
    )
    after = sha256_file(path)

    assert master.audit.status == "ok"
    assert master.audit.rows == 2
    assert master.audit.columns == 41
    assert master.audit.formula_count == 0
    assert master.audit.read_only is True
    assert master.audit.writes_master is False
    assert master.audit.hash_unchanged_after_read is True
    assert before == after


def test_rejects_wrong_hash(
    tmp_path: Path,
) -> None:
    path = tmp_path / "trades_master.xlsx"
    _write_workbook(path, [_row(1)])

    contract = replace(
        _contract_for(path, 1),
        expected_sha256="0" * 64,
    )

    with pytest.raises(
        OfficialMasterValidationError,
        match="master_sha256_mismatch",
    ):
        load_official_trades_master(
            path,
            contract=contract,
        )


def test_rejects_schema_order_drift(
    tmp_path: Path,
) -> None:
    path = tmp_path / "trades_master.xlsx"
    _write_workbook(path, [_row(1)])

    workbook = load_workbook(path)
    worksheet = workbook["TRADES"]
    worksheet["A1"] = "unexpected_header"
    workbook.save(path)
    workbook.close()

    with pytest.raises(
        OfficialMasterValidationError,
        match="schema_column_order_mismatch",
    ):
        load_official_trades_master(
            path,
            contract=_contract_for(path, 1),
        )


def test_rejects_economic_identity_drift(
    tmp_path: Path,
) -> None:
    path = tmp_path / "trades_master.xlsx"
    row = _row(1)
    row["economic_net_pnl"] = 999.0
    _write_workbook(path, [row])

    with pytest.raises(
        OfficialMasterValidationError,
        match="economic_identity_drift",
    ):
        load_official_trades_master(
            path,
            contract=_contract_for(path, 1),
        )


def test_rejects_formulas(
    tmp_path: Path,
) -> None:
    path = tmp_path / "trades_master.xlsx"
    _write_workbook(path, [_row(1)])

    workbook = load_workbook(path)
    worksheet = workbook["TRADES"]
    worksheet["F2"] = "=1+1"
    workbook.save(path)
    workbook.close()

    with pytest.raises(
        OfficialMasterValidationError,
        match="formula_count_mismatch",
    ):
        load_official_trades_master(
            path,
            contract=_contract_for(path, 1),
        )


def test_rejects_invalid_order_id(
    tmp_path: Path,
) -> None:
    path = tmp_path / "trades_master.xlsx"
    _write_workbook(
        path,
        [
            _row(
                1,
                order_id="INVALID",
            )
        ],
    )

    with pytest.raises(
        OfficialMasterValidationError,
        match="invalid_or_missing_order_id",
    ):
        load_official_trades_master(
            path,
            contract=_contract_for(path, 1),
        )


def test_rejects_non_contiguous_trade_sequence(
    tmp_path: Path,
) -> None:
    path = tmp_path / "trades_master.xlsx"
    _write_workbook(
        path,
        [
            _row(1),
            _row(3),
        ],
    )

    with pytest.raises(
        OfficialMasterValidationError,
        match="trade_sequence_not_contiguous",
    ):
        load_official_trades_master(
            path,
            contract=_contract_for(path, 2),
        )
